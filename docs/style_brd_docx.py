# Styles the pandoc export of BRD.md:  pandoc BRD.md -f gfm -t docx -o BRD.docx && python3 style_brd_docx.py BRD.docx
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
path=sys.argv[1]
d=Document(path)
sec=d.sections[0]; sec.left_margin=sec.right_margin=Inches(0.9)
sizes={'Heading 1':18,'Heading 2':14,'Heading 3':11}
for st in d.styles:
    try:
        if st.type!=1: continue
        st.font.name='Arial'
        rpr=st.element.get_or_add_rPr(); rf=rpr.find(qn('w:rFonts'))
        if rf is None: rf=OxmlElement('w:rFonts'); rpr.append(rf)
        for a in ('w:ascii','w:hAnsi','w:cs','w:eastAsia'): rf.set(qn(a),'Arial')
        for a in ('w:asciiTheme','w:hAnsiTheme','w:cstheme','w:eastAsiaTheme'):
            if rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
        if st.name in sizes: st.font.size=Pt(sizes[st.name]); st.font.color.rgb=RGBColor(0x14,0x30,0x3A)
        if st.name in ('Normal','Body Text','First Paragraph','Compact'): st.font.size=Pt(10)
    except Exception as e: print('skip',st.name,e)
def shade(cell,color):
    tcPr=cell._tc.get_or_add_tcPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),color); tcPr.append(sh)
total=6.7
for t in d.tables:
    tblPr=t._tbl.tblPr
    for old in tblPr.findall(qn('w:tblW')): tblPr.remove(old)
    w=OxmlElement('w:tblW'); w.set(qn('w:w'),str(int(total*1440))); w.set(qn('w:type'),'dxa'); tblPr.append(w)
    lay=OxmlElement('w:tblLayout'); lay.set(qn('w:type'),'fixed'); tblPr.append(lay)
    b=OxmlElement('w:tblBorders')
    for e in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{e}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),'B9C8C5'); b.append(el)
    tblPr.append(b)
    n=len(t.columns)
    weights=[max(4,min(70,max(len(r.cells[i].text) for r in t.rows))) for i in range(n)]
    s=sum(weights); widths=[max(0.75,total*x/s) for x in weights]
    k=total/sum(widths); widths=[x*k for x in widths]
    for i,gc in enumerate(t._tbl.tblGrid.findall(qn('w:gridCol'))): gc.set(qn('w:w'),str(int(widths[i]*1440)))
    for row in t.rows:
        for i,c in enumerate(row.cells):
            c.width=Inches(widths[i])
            for p in c.paragraphs:
                for r in p.runs: r.font.size=Pt(9); r.font.name='Arial'
    for c in t.rows[0].cells:
        shade(c,'14303A')
        for p in c.paragraphs:
            for r in p.runs: r.font.color.rgb=RGBColor(255,255,255); r.font.bold=True
d.save(path)
