"""Builds test-cases.xlsx from results.json (written by prototype/tests/e2e.mjs)."""
import json
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).parent
data = json.loads((HERE / "results.json").read_text())
results = {r["id"]: r for r in data["results"]}
cases = data["cases"]
run_date = "16/10/2026"

F = "Arial"
thin = Side(style="thin", color="B9C8C5")
border = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="14303A")
head_font = Font(name=F, bold=True, color="FFFFFF", size=10)
body = Font(name=F, size=10)
wrap = Alignment(wrap_text=True, vertical="top")

wb = Workbook()

# ---- Test Cases sheet ----
ws = wb.active
ws.title = "Test Cases"
ws["A1"] = "Hospital Appointment Booking System — Test Case Sheet"
ws["A1"].font = Font(name=F, bold=True, size=14, color="14303A")
ws["A2"] = (f"Executed {run_date} against the clickable prototype (prototype/dist) with the browser clock fixed at "
            "Wed 14 Oct 2026 08:00 (Europe/London) unless the preconditions say otherwise. "
            "Automated with Playwright: prototype/tests/e2e.mjs. Status column has a Pass / Fail / Not Executed drop-down.")
ws["A2"].font = Font(name=F, size=9, italic=True, color="3E5862")
ws.merge_cells("A2:M2")
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[2].height = 30

headers = ["Test case ID", "Test area", "Requirement / rule", "Priority", "Test scenario", "Preconditions",
           "Test steps", "Test data", "Expected results", "Actual results", "Test status", "Executed on", "Executed by"]
widths = [10, 18, 16, 9, 34, 34, 48, 30, 46, 50, 12, 12, 16]
HR = 4
for i, (h, w) in enumerate(zip(headers, widths), start=1):
    c = ws.cell(row=HR, column=i, value=h)
    c.font, c.fill, c.border = head_font, head_fill, border
    c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.column_dimensions[c.column_letter].width = w
ws.row_dimensions[HR].height = 28

for r, tc in enumerate(cases, start=HR + 1):
    res = results.get(tc["id"], {})
    steps = "\n".join(f"{n}. {s}" for n, s in enumerate(tc["steps"], 1))
    row = [tc["id"], tc["area"], tc["req"], tc["priority"], tc["scenario"], tc["pre"], steps, tc["data"],
           tc["expected"], res.get("actual", ""), res.get("status", "Not Executed"),
           run_date if res else "", "Automated (Playwright)" if res else ""]
    for c, v in enumerate(row, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font, cell.alignment, cell.border = body, wrap, border
    ws.cell(row=r, column=1).font = Font(name=F, size=10, bold=True)
    ws.cell(row=r, column=11).alignment = Alignment(horizontal="center", vertical="top")
    ws.cell(row=r, column=11).font = Font(name=F, size=10, bold=True)
last = HR + len(cases)

dv = DataValidation(type="list", formula1='"Pass,Fail,Not Executed"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f"K{HR+1}:K{last}")
rng = f"K{HR+1}:K{last}"
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Pass"'], fill=PatternFill("solid", fgColor="E3F2ED"), font=Font(color="17735F", bold=True)))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Fail"'], fill=PatternFill("solid", fgColor="FBECEB"), font=Font(color="B3261E", bold=True)))
ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"Not Executed"'], fill=PatternFill("solid", fgColor="FDF3DC"), font=Font(color="8A5A00", bold=True)))
ws.freeze_panes = "B5"
ws.auto_filter.ref = f"A{HR}:M{last}"
ws.page_setup.orientation = "landscape"
ws.page_setup.fitToWidth = 1
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToHeight = 0
ws.print_title_rows = f"{HR}:{HR}"

# ---- Summary sheet ----
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "Test execution summary"
sm["A1"].font = Font(name=F, bold=True, size=14, color="14303A")
sm["A2"] = f"Prototype build tested on {run_date}. Counts are formulas over the Test Cases sheet and update if a status is changed."
sm["A2"].font = Font(name=F, size=9, italic=True, color="3E5862")
TC = "'Test Cases'"
st_rng = f"{TC}!$K${HR+1}:$K${last}"
ar_rng = f"{TC}!$B${HR+1}:$B${last}"
for c, h in enumerate(["Status", "Count", "% of total"], start=1):
    cell = sm.cell(row=4, column=c, value=h); cell.font, cell.fill, cell.border = head_font, head_fill, border
for i, s in enumerate(["Pass", "Fail", "Not Executed"]):
    r = 5 + i
    sm.cell(row=r, column=1, value=s)
    sm.cell(row=r, column=2, value=f'=COUNTIF({st_rng},A{r})')
    sm.cell(row=r, column=3, value=f'=IFERROR(B{r}/$B$8,0)').number_format = "0.0%"
sm.cell(row=8, column=1, value="Total")
sm.cell(row=8, column=2, value="=SUM(B5:B7)")
sm.cell(row=8, column=3, value="=SUM(C5:C7)").number_format = "0.0%"
for r in range(5, 9):
    for c in range(1, 4):
        cell = sm.cell(row=r, column=c); cell.border = border
        cell.font = Font(name=F, size=10, bold=(r == 8))

areas = list(dict.fromkeys(tc["area"] for tc in cases))
sm.cell(row=10, column=1, value="By test area").font = Font(name=F, bold=True, size=11)
for c, h in enumerate(["Test area", "Total", "Pass", "Fail", "Not Executed"], start=1):
    cell = sm.cell(row=11, column=c, value=h); cell.font, cell.fill, cell.border = head_font, head_fill, border
for i, a in enumerate(areas):
    r = 12 + i
    sm.cell(row=r, column=1, value=a)
    sm.cell(row=r, column=2, value=f'=COUNTIF({ar_rng},A{r})')
    for c, s in zip((3, 4, 5), ("Pass", "Fail", "Not Executed")):
        sm.cell(row=r, column=c, value=f'=COUNTIFS({ar_rng},$A{r},{st_rng},"{s}")')
    for c in range(1, 6):
        sm.cell(row=r, column=c).font = body; sm.cell(row=r, column=c).border = border
for col, w in zip("ABCDE", (32, 10, 10, 10, 14)):
    sm.column_dimensions[col].width = w

n = 12 + len(areas) + 1
notes = [
    "Scope: the eight areas required by the brief (TC-01 – TC-32) plus admin, validation, booking-limit and responsive edge cases (TC-33 – TC-38).",
    "Environment: Chromium (Playwright), desktop 1280×860 and mobile 390×844, fresh browser storage per test case.",
    "Business rules are also unit-tested in prototype/tests/rules.test.js (npm test).",
    "Limitation: the prototype has no backend, so concurrency (TC-06) is simulated with an in-page demo control; production needs a server-side unique constraint (NFR-03).",
]
sm.cell(row=n, column=1, value="Notes").font = Font(name=F, bold=True, size=11)
for i, t in enumerate(notes, start=1):
    c = sm.cell(row=n + i, column=1, value=f"• {t}"); c.font = body
    sm.merge_cells(start_row=n + i, start_column=1, end_row=n + i, end_column=5)
    c.alignment = Alignment(wrap_text=True, vertical="top"); sm.row_dimensions[n + i].height = 28

# ---- Traceability sheet ----
tr = wb.create_sheet("Traceability")
tr["A1"] = "Requirements traceability: business rules → test cases"
tr["A1"].font = Font(name=F, bold=True, size=14, color="14303A")
rules = [
    ("BR-01", "Only registered, logged-in patients can book", "TC-02, TC-21, TC-23, TC-25"),
    ("BR-02", "Patient cannot book an occupied slot", "TC-04, TC-05, TC-06, TC-08"),
    ("BR-03", "Doctor cannot have two appointments in one slot", "TC-08"),
    ("BR-04", "Must select a valid available date and time", "TC-11, TC-12"),
    ("BR-05", "Cancelled appointment releases the slot", "TC-13, TC-14, TC-16, TC-26"),
    ("BR-06", "Reschedule only to an available slot", "TC-17, TC-18, TC-20"),
    ("BR-07", "Doctors see/manage only their own appointments", "TC-24"),
    ("BR-08", "Admins manage doctors, patients, appointments", "TC-20, TC-26, TC-33, TC-34, TC-35"),
    ("BR-09", "30-min slots, ≤30 days ahead, no past slots", "TC-09, TC-10, TC-12, TC-30"),
    ("BR-10", "No online change within 2 hours of start", "TC-15, TC-19"),
    ("BR-11", "Patient cannot hold two appointments at the same time", "TC-07"),
    ("BR-12", "Booked slot cannot be removed from availability", "TC-29"),
    ("BR-13", "Complete only after start time", "TC-31"),
    ("BR-14", "Reschedule keeps the same doctor", "TC-17"),
    ("BR-15", "Inactive doctors hidden and not bookable", "TC-33"),
    ("BR-16", "Max 3 upcoming appointments per patient", "TC-37"),
    ("FR-41/42", "Doctor publishes and removes availability", "TC-27, TC-28, TC-32"),
    ("NFR-01", "Responsive UI", "TC-38"),
]
for c, h in enumerate(["Rule / req.", "Description", "Test cases"], start=1):
    cell = tr.cell(row=3, column=c, value=h); cell.font, cell.fill, cell.border = head_font, head_fill, border
for i, row in enumerate(rules, start=4):
    for c, v in enumerate(row, start=1):
        cell = tr.cell(row=i, column=c, value=v); cell.font, cell.border = body, border
        cell.alignment = Alignment(wrap_text=True, vertical="top")
for col, w in zip("ABC", (12, 52, 40)):
    tr.column_dimensions[col].width = w

out = HERE / "test-cases.xlsx"
wb.save(out)
print("saved", out, len(cases), "cases")
